#!/usr/bin/env python3
"""How many households own less than nothing, and which ones.

The atlas already reports that in five countries the poorest half owns less than
nothing. That is a share of a national total: add up what half the country holds and
the sum comes out negative. It says nothing about how many people are in that
position, and a share can be dragged under by a small number of households deeply in
debt. This is the headcount beside it.

Reads the four HFCS statistical-table workbooks in data/raw/hfcs/ and writes
data/hfcs_negative.json.

It is a different measure from the wealth shares and must never be drawn on the same
axis as them:

  the shares      WID, per adult, an aggregate share of a national total, 213
                  countries, annual, mostly imputed
  this            ECB HFCS, per household, a count of households below zero, euro area
                  plus a few neighbours, four survey waves, measured

Table F3 of each wave, which carries the total and four breakdowns. Three are kept:

  all             every household
  housing         outright owners, owners with a mortgage, renters and others
  age             the age of the reference person, in six bands

The income breakdown is dropped. It is a headcount of negative net *wealth* cut by
*income*, and the atlas's whole argument is that those two are different things, so
the cut invites exactly the conflation the rest of the page works to prevent.

Three things about the source that the app has to repeat rather than hide:

  N and M         too few observations, and missing. Stored as null, never as zero.
  "< 0.1"         the source's way of saying below a tenth of a per cent. Stored at
                  0.1, its stated ceiling, so a reading built on it can never claim
                  the number is smaller than the source will support.
  the wave year   waves are named 2010, 2014, 2017 and 2021 and the fieldwork behind
                  each ran in different years in different countries. The year here
                  is the wave's name, not a reference date, and the app says so.

The country set changes between waves, so the euro-area aggregate is not the same
aggregate across waves either. The app says that too.

Usage: pip install openpyxl && python3 build_hfcs.py
"""
import json
import os
import re

try:
    import openpyxl
except ImportError:
    raise SystemExit('needs openpyxl: pip install openpyxl')

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'data')
RAW = os.path.join(DATA, 'raw', 'hfcs')

# the rows of table F3 worth keeping, by the label in its second column
ROWS = {
    'ALL': 'all',
    'Owners - outright': 'own_outright',
    'Owners - with mortgage': 'own_mortgage',
    'Renters / other': 'renting',
    '16-34': 'age_16_34',
    '35-44': 'age_35_44',
    '45-54': 'age_45_54',
    '55-64': 'age_55_64',
    '65-74': 'age_65_74',
    '75+': 'age_75_plus',
}
# HFCS reports two-letter codes; the atlas keys on ISO3
ISO3 = {
    'BE': 'BEL', 'CZ': 'CZE', 'DE': 'DEU', 'EE': 'EST', 'IE': 'IRL', 'GR': 'GRC',
    'ES': 'ESP', 'FR': 'FRA', 'HR': 'HRV', 'IT': 'ITA', 'CY': 'CYP', 'LV': 'LVA',
    'LT': 'LTU', 'LU': 'LUX', 'HU': 'HUN', 'MT': 'MLT', 'NL': 'NLD', 'AT': 'AUT',
    'PL': 'POL', 'PT': 'PRT', 'SI': 'SVN', 'SK': 'SVK', 'FI': 'FIN',
}


def value(cell):
    """A percentage, or None where the source says it cannot give one."""
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        return round(float(cell), 2)
    t = str(cell).strip()
    if t in ('N', 'M', ''):
        return None
    if t.startswith('<'):
        m = re.search(r'([\d.]+)', t)
        return round(float(m.group(1)), 2) if m else None   # the stated ceiling
    try:
        return round(float(t), 2)
    except ValueError:
        return None


def read_wave(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = [s for s in wb.sheetnames if s.startswith('F3')]
    if not sheet:
        wb.close()
        return None
    ws = wb[sheet[0]]
    rows = list(ws.iter_rows(max_row=60, max_col=40, values_only=True))
    wb.close()
    # row 4 carries the country codes; every later row lines up on the same columns
    head = rows[3]
    cols = {}
    for i, c in enumerate(head):
        if not c:
            continue
        name = str(c).replace('\n', ' ').strip()
        if name.lower() == 'euro area':
            cols[i] = 'EA'
        elif name in ISO3:
            cols[i] = ISO3[name]
    out = {}
    for r in rows:
        label = str(r[1]).strip() if r[1] else ''
        key = ROWS.get(label)
        if not key:
            continue
        for i, iso in cols.items():
            v = value(r[i])
            if v is None:
                continue
            out.setdefault(iso, {})[key] = v
    return out


def main():
    if not os.path.isdir(RAW):
        raise SystemExit('missing %s\n%s' % (RAW, __doc__))
    waves = {}
    for fn in sorted(os.listdir(RAW)):
        if not fn.endswith('.xlsx'):
            continue
        m = re.search(r'(20\d{2})', fn)
        if not m:
            continue
        got = read_wave(os.path.join(RAW, fn))
        if got:
            waves[int(m.group(1))] = got

    years = sorted(waves)
    out = {
        'unit': 'per cent of households whose debts exceed their assets',
        'source': 'ECB Household Finance and Consumption Survey, statistical table F3, '
                  'waves 2010, 2014, 2017 and 2021.',
        'waves': years,
        'cuts': list(dict.fromkeys(ROWS.values())),
        'by': {},
    }
    every = sorted({iso for w in waves.values() for iso in w})
    for iso in every:
        block = {}
        for cut in out['cuts']:
            series = [waves[y].get(iso, {}).get(cut) for y in years]
            if any(v is not None for v in series):
                block[cut] = series
        if block:
            out['by'][iso] = block

    path = os.path.join(DATA, 'hfcs_negative.json')
    with open(path, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, separators=(',', ':'), sort_keys=True)

    print('waves      %s' % ', '.join(str(y) for y in years))
    print('countries  %d, plus the euro area aggregate'
          % len([i for i in out['by'] if i != 'EA']))
    print('wrote %s (%d KB)' % (os.path.relpath(path, HERE), os.path.getsize(path) // 1024))

    last = len(years) - 1
    ea = out['by'].get('EA', {})
    print('\neuro area, wave %d:' % years[last])
    for cut in out['cuts']:
        v = ea.get(cut, [None] * len(years))[last]
        if v is not None:
            print('   %-14s %5.1f%%' % (cut, v))
    print('\nhighest in the latest wave:')
    rank = [(b['all'][last], i) for i, b in out['by'].items()
            if i != 'EA' and b.get('all') and b['all'][last] is not None]
    for v, i in sorted(rank, reverse=True)[:6]:
        print('   %s %5.1f%%' % (i, v))


if __name__ == '__main__':
    main()
